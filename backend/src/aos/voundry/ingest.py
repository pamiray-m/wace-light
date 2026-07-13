"""
WACE Smart Workspace — file ingestion.

Parses a dropped Excel / Word / PDF / text file IN MEMORY, scrubs it with SAIb
BEFORE any of it reaches an agent/LLM, receipts the ingest, and returns the
scrubbed text. The file itself is never persisted (ephemeral) — the operator can
save the resulting analysis to the Flight Recorder if they want to keep it.
"""

from __future__ import annotations

import io

from src.aos.voundry.contracts import WorkUnit
from src.aos.voundry.governance import voundry_audit
from src.aos.voundry.persistence.repository import voundry_repo

_MAX_BYTES = 12 * 1024 * 1024
_MAX_TEXT = 20000


class IngestError(Exception):
    pass


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Return (kind, extracted_text) for a supported file — in-memory only."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets[:10]:
            lines.append(f"# Sheet: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 500:
                    lines.append("… (truncated)")
                    break
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "spreadsheet", "\n".join(lines)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        return "document", "\n".join(p.text for p in d.paragraphs if p.text.strip())
    if name.endswith(".pdf"):
        import pypdf
        r = pypdf.PdfReader(io.BytesIO(data))
        return "pdf", "\n".join((p.extract_text() or "") for p in r.pages[:50])
    if name.endswith((".txt", ".csv", ".md", ".log", ".json", ".yaml", ".yml", ".ini", ".conf")):
        return "text", data.decode("utf-8", errors="replace")
    ext = name.rsplit(".", 1)[-1] if "." in name else "?"
    raise IngestError(f"Unsupported file type '.{ext}'. Try Excel, Word, PDF, or a text file.")


def preview_redaction(text: str) -> dict:
    """Show what SAIb WILL mask before any content reaches an agent.

    Returns the masked spans (type + offsets) and a per-type tally so the UI can
    highlight, in real time, exactly what gets scrubbed — the single most
    trust-building thing a compliance officer can watch. Read-only, no LLM call.
    """
    text = text or ""
    try:
        from src.saib.detectors import detector_registry
        entities = detector_registry.scan(text)
    except Exception:  # noqa: BLE001 — preview must never raise
        entities = []
    spans = [{"type": e.type, "start": e.start, "end": e.end} for e in entities
             if isinstance(getattr(e, "start", None), int) and isinstance(getattr(e, "end", None), int)]
    by_type: dict[str, int] = {}
    for e in entities:
        by_type[e.type] = by_type.get(e.type, 0) + 1
    masked, _, _ = _scrub(text)
    return {"count": len(spans), "spans": spans, "by_type": by_type, "masked": masked[:_MAX_TEXT]}


def _scrub(text: str) -> tuple[str, int, bool]:
    """Return (scrubbed_text, masked_entity_count, ok).

    ``ok`` is False when the SAIb guard is unavailable or errors. Ingest is a
    core data-safety promise ("scrubbed BEFORE any of it reaches an agent/LLM"),
    so callers MUST fail closed on ``ok is False`` rather than pass raw content.
    """
    if not text:
        return text, 0, True
    try:
        from src.saib.guard import saib_guard
        result = saib_guard.process(text, "")
        masked = getattr(result, "safe_prompt", None)
        entities = len(getattr(result, "entities", []) or [])
        if masked is not None:
            return masked, entities, True
    except Exception:  # noqa: BLE001 — a guard failure must fail CLOSED, not open
        pass
    return text, 0, False


def ingest_file(contributor_id: str, work_unit_id: str, filename: str, data: bytes,
                repo=voundry_repo, audit=voundry_audit) -> dict:
    w = repo.get_work_unit(work_unit_id)
    if w is None:
        raise IngestError("Work unit not found.")
    if WorkUnit(**w).assigned_to != contributor_id:
        raise IngestError("Not your assignment.")
    if not data:
        raise IngestError("Empty file.")
    if len(data) > _MAX_BYTES:
        raise IngestError("File too large (max 12 MB).")
    kind, text = extract_text(filename, data)
    scrubbed, masked, scrub_ok = _scrub(text)
    if not scrub_ok:
        # Fail closed: never return unscrubbed content when the guard is down.
        raise IngestError("Content-safety scrub is unavailable — ingest refused. Try again shortly.")
    scrubbed = scrubbed[:_MAX_TEXT]
    # For spreadsheets, hand the UI a structured table parsed from the SAME
    # scrubbed text (so the rendered grid is SAIb-safe, never the raw cells).
    table = _table_from_spreadsheet(scrubbed) if kind == "spreadsheet" else None
    audit.append(
        actor_id=contributor_id, actor_type="human", action="file.ingested",
        resource_type="work_unit", resource_id=work_unit_id,
        detail=f"ingested {filename} ({kind}, {masked} masked)",
        metadata={"filename": (filename or "")[:120], "kind": kind, "masked_entities": masked, "chars": len(scrubbed)},
    )
    out = {"name": filename, "kind": kind, "text": scrubbed, "masked_entities": masked, "chars": len(scrubbed)}
    if table:
        out["table"] = table
    return out


_MAX_TABLE_ROWS = 100
_MAX_TABLE_COLS = 25


def _table_from_spreadsheet(scrubbed: str) -> list[list[str]]:
    """Split the scrubbed spreadsheet text back into a row/cell grid for display.

    Sheet headers ('# Sheet: …') and the truncation marker are dropped; each data
    line is a ' | '-joined row (see extract_text), so we split on that delimiter.
    """
    rows: list[list[str]] = []
    for line in scrubbed.split("\n"):
        if not line.strip() or line.startswith("# Sheet:") or line.startswith("…"):
            continue
        rows.append(line.split(" | ")[:_MAX_TABLE_COLS])
        if len(rows) >= _MAX_TABLE_ROWS:
            break
    return rows
